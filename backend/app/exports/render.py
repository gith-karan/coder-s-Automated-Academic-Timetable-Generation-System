"""Timetable exports. All formats are rendered from one view model (sheets of day × period cells)."""
from __future__ import annotations

import csv
import io
from xml.sax.saxutils import escape
from datetime import date, datetime, timedelta

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(v: str) -> str:
    """Neutralise spreadsheet formula injection (CSV/Excel)."""
    return "'" + v if v and v.startswith(_FORMULA_PREFIXES) else v


def cell_text(e: dict) -> str:
    parts = [f'{e["subject"].get("short_name") or e["subject"]["code"]} · {e["subject"]["code"]}']
    if e.get("batch"):
        parts[0] += f" ({e['batch']['name']})"
    if e.get("room"):
        parts.append(e["room"]["code"])
    if e.get("faculty"):
        parts.append(", ".join(f["name"] for f in e["faculty"]))
    return "\n".join(parts)


def build_sheets(view: dict, groups: list[tuple[str, str]], key: str) -> list[dict]:
    """groups: [(id, title)]; key: entry field to split sheets by ("group"/"faculty"/"room")."""
    grid = view["grid"]
    days = grid["days"]
    max_p = max((len(d["periods"]) for d in days), default=0)
    sheets = []
    for gid, title in groups:
        cells: dict[tuple[int, int], list[str]] = {}
        for e in view["entries"]:
            match = (e["group"]["id"] == gid if key == "group" else
                     any(f["id"] == gid for f in e["faculty"]) if key == "faculty" else
                     (e["room"] or {}).get("id") == gid)
            if not match:
                continue
            for i in range(e["length"]):
                cells.setdefault((e["day"], e["start_period"] + i), []).append(cell_text(e))
        rows = []
        for pi in range(max_p):
            label = next((f"{d['periods'][pi]['start']}–{d['periods'][pi]['end']}" for d in days
                          if pi < len(d["periods"])), f"P{pi + 1}")
            rows.append([label] + [" / ".join(cells.get((d["day"], pi), [])) for d in days])
        sheets.append({"title": title, "header": ["Time"] + [d["name"] for d in days], "rows": rows})
    return sheets


def to_csv(sheets: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    for sh in sheets:
        w.writerow([safe_cell(sh["title"])])
        w.writerow(sh["header"])
        for r in sh["rows"]:
            w.writerow([safe_cell(c.replace("\n", " | ")) for c in r])
        w.writerow([])
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(sheets: list[dict], title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    wb.remove(wb.active)
    maroon = PatternFill("solid", fgColor="7A1F3D")
    thin = Side(style="thin", color="D9D9D9")
    for sh in sheets:
        name = "".join(c for c in sh["title"] if c not in r"[]:*?/\\")[:31] or "Sheet"
        ws = wb.create_sheet(name)
        ws.append([safe_cell(f"{title} — {sh['title']}")])
        ws["A1"].font = Font(bold=True, size=13, color="7A1F3D")
        ws.append(sh["header"])
        for c in ws[2]:
            c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), maroon, Alignment(horizontal="center")
        for r in sh["rows"]:
            ws.append([safe_cell(c) for c in r])
        for row in ws.iter_rows(min_row=3):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        ws.column_dimensions["A"].width = 14
        for col in "BCDEFGH":
            ws.column_dimensions[col].width = 24
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_pdf(sheets: list[dict], title: str, subtitle: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=title)
    styles = getSampleStyleSheet()
    cell_style = styles["BodyText"].clone("cell", fontSize=7.5, leading=9)
    maroon = colors.HexColor("#7A1F3D")
    story = []
    for i, sh in enumerate(sheets):
        story.append(Paragraph("<font color='#7A1F3D'><b>National Forensic Sciences University</b></font>", styles["Title"]))
        story.append(Paragraph(f"{escape(title)} — <b>{escape(sh['title'])}</b>", styles["Heading3"]))
        story.append(Paragraph(escape(subtitle), styles["Normal"]))
        story.append(Spacer(1, 4 * mm))
        data = [sh["header"]] + [[r[0]] + [Paragraph(escape(c).replace("\n", "<br/>"), cell_style)
                                           for c in r[1:]] for r in sh["rows"]]
        ncols = len(sh["header"])
        width = doc.width
        t = Table(data, colWidths=[22 * mm] + [(width - 22 * mm) / (ncols - 1)] * (ncols - 1), repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), maroon), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 8), ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D0D0D0")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FAF6F7")]),
        ]))
        story.append(t)
        if i < len(sheets) - 1:
            story.append(PageBreak())
    if not story:
        story.append(Paragraph("No sessions to show.", styles["Normal"]))
    doc.build(story)
    return buf.getvalue()


def to_ics(view: dict, title: str, term_start: date | None = None, weeks: int = 16) -> bytes:
    from icalendar import Calendar, Event

    cal = Calendar()
    cal.add("prodid", "-//NFSU//Timetable System//EN")
    cal.add("version", "2.0")
    cal.add("x-wr-calname", title)
    start = term_start or (date.today() + timedelta(days=(7 - date.today().weekday()) % 7 or 7))
    monday = start - timedelta(days=start.weekday())
    days = {d["day"]: d for d in view["grid"]["days"]}
    for e in view["entries"]:
        d = days.get(e["day"])
        if not d:
            continue
        p0, p1 = d["periods"][e["start_period"]], d["periods"][e["start_period"] + e["length"] - 1]
        day = monday + timedelta(days=e["day"])
        sh, sm = map(int, p0["start"].split(":"))
        eh, em = map(int, p1["end"].split(":"))
        ev = Event()
        ev.add("uid", f"{e['session_key']}@timetable.nfsu.ac.in")
        ev.add("summary", f"{e['subject']['code']} {e['subject']['name']}" + (f" ({e['batch']['name']})" if e.get("batch") else ""))
        ev.add("dtstart", datetime(day.year, day.month, day.day, sh, sm))
        ev.add("dtend", datetime(day.year, day.month, day.day, eh, em))
        ev.add("rrule", {"freq": "weekly", "count": weeks})
        if e.get("room"):
            ev.add("location", f"{e['room']['code']} — {e['room']['name']}")
        ev.add("description", f"{e['group']['label']}\nFaculty: {', '.join(f['name'] for f in e['faculty'])}")
        cal.add_component(ev)
    return cal.to_ical()
